import pptx

prs = pptx.Presentation('docs/Ananta_Submission_Deck.pptx')

# Slide 7 (index 6): Architecture
s7 = prs.slides[6]
for s in s7.shapes:
    if s.name == 'Rounded Rectangle 116' and s.has_text_frame:
        s.text_frame.text = 'Guided session + tool dispatch\nServer owns the board; 8 live tools in guided chalkboard (13 in standard)'

# Slide 8 (index 7): Technologies
s8 = prs.slides[7]
for s in s8.shapes:
    if s.name == 'Rectangle 121' and s.has_text_frame:
        s.text_frame.text = 'Real-time two-way voice, voice activity detection, 8 function-calling tools in default guided chalkboard (13 in standard)'
    elif s.name == 'Rectangle 125' and s.has_text_frame:
        s.text_frame.text = 'Per-answer reasoning diagnosis with 6.5s budget (measured p50: 3.15s, p95: 4.82s)'
    elif s.name == 'Rectangle 127' and s.has_text_frame:
        s.text_frame.text = 'Fallback for diagnosis and generation (resilience)'

# Slide 10 (index 9): Snapshots
s10 = prs.slides[9]
for s in s10.shapes:
    if s.name == 'Rectangle 144' and s.has_text_frame:
        s.text_frame.text = 'Live app: https://ananta-ivczdbgbja-as.a.run.app/  ·  no sign-in: pick "Maya (Simulated)" or "Leo (Simulated)", Mathematics, Start Lesson'

# Slide 11 (index 10): Performance & Benchmarking
s11 = prs.slides[10]
for s in s11.shapes:
    if s.name == 'Rectangle 143' and s.has_text_frame:
        s.text_frame.text = 'Diagnosis accuracy on 45 labelled exchanges (run 2026-10-17 with gemini-3.6-flash).'
    elif s.name == 'Rectangle 151' and s.has_text_frame:
        s.text_frame.text = '10/10'
    elif s.name == 'Rectangle 155' and s.has_text_frame:
        s.text_frame.text = '10/10'
    elif s.name == 'Rectangle 156' and s.has_text_frame:
        s.text_frame.text = '96%\nexact class, 43/45'
    elif s.name == 'Rectangle 160' and s.has_text_frame:
        s.text_frame.text = (
            'Honest numbers\n'
            'Latency per diagnosis in this run: p50 3.15 s, p95 4.82 s, max 5.79 s, mean 3.32 s (n=45) — well under our 6.5s budget!\n'
            'Flags are 100% precise: 10 of 10 flags were right (plain prompt: 10 of 11, 91%).\n'
            'Adaptation proof: npm run test:adaptation (10/10 automated checks passing).'
        )
    elif s.name == 'Rectangle 161' and s.has_text_frame:
        s.text_frame.text = (
            'How to read this\n'
            'Run on gemini-3.6-flash (2026-10-17). 45 hand-written cases across 5 concepts with Wilson 95% intervals.\n'
            'Only 10 hidden-misconception cases: 95% CI for 10/10 is 72–100%. Tests the single-answer assessor.\n'
            'Also verified: app shell loads in 102 ms; learner data server-only; test suite passing for quality gates, BKT, deterministic compiler, and parent evidence replay.'
        )

# Slide 13 (index 12): Links & Evaluator guide
s13 = prs.slides[12]
for s in s13.shapes:
    if s.name == 'Rectangle 159' and s.has_text_frame:
        s.text_frame.text = (
            'For evaluators: test in under a minute\n'
            '1.  Open the product link. No sign-in is needed (demo mode).\n'
            '2.  Click "Maya (Simulated)" or "Leo (Simulated)" (Grade 4 Mathematics, Money and decimal place value). '
            'Notice Ananta opens with a place-value diagram for Maya, but with a worked example and avoids diagrams for Leo.\n'
            '3.  Try asking "Can you show me a picture?" or answering "10 coins" to see the diagnostic assessment and BKT update in real time.\n'
            '4.  Open the Parent Portal from the start page to see Ananta\'s plan for the next lesson and click "See Leo\'s answer" to inspect the child\'s exact words.\n'
            'First load may take 10–20 s. Chrome or Firefox recommended.'
        )

output_path = 'docs/Submission_Template_AI_Builder_Cup_Filled.pptx'
prs.save(output_path)
print(f'Successfully created {output_path}!')

try:
    prs.save('docs/Ananta_Submission_Deck.pptx')
    print('Successfully updated docs/Ananta_Submission_Deck.pptx!')
except PermissionError:
    print('Notice: docs/Ananta_Submission_Deck.pptx is currently open in PowerPoint. Saved to docs/Submission_Template_AI_Builder_Cup_Filled.pptx instead.')

